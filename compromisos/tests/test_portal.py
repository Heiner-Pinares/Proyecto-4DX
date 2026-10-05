from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.conf import settings
from django.contrib.auth.models import Group, User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from docx import Document
from openpyxl import load_workbook

from compromisos.indicators import metricas
from compromisos.models import Compromiso
from compromisos.reports import excel, word
from compromisos.selectors import base, filtrar
from compromisos.services import fecha_objetivo, guardar, operar, situacion

pytestmark = pytest.mark.django_db


@pytest.fixture
def users():
    call_command("init_portal")
    result = {}
    for role in ["Administrador", "Editor", "Consulta"]:
        user = User.objects.create_user(role, password="Test-2026-safe")
        user.groups.add(Group.objects.get(name=role))
        result[role] = user
    return result


@pytest.fixture
def data():
    return dict(
        tema="Proyectos",
        iniciativa="Proyectos Internos",
        tarea="Preparar formatos",
        responsable_pyp="HP",
        status="EC",
        fecha_de_compromiso=date(2026, 9, 1),
        fecha_de_vencimiento=date(2026, 9, 10),
        meta=Decimal("95"),
        puntaje=Decimal("100"),
    )


@pytest.fixture
def c(data, users):
    return guardar(data, users["Editor"])


def test_creation_and_edit(c, users):
    assert c.mes == "Septiembre 2026"
    assert c.historial.get().accion == "CREADO"
    updated = guardar({"meta": Decimal("110")}, users["Editor"], c.pk)
    assert updated.meta == 110
    assert "95" in updated.historial.first().descripcion
    assert updated.updated_by == "Editor"
    assert updated.historial.count() == 2


def test_soft_delete_restore_purge(c, users):
    operar(c.pk, "eliminar", {}, users["Administrador"])
    assert not Compromiso.objects.filter(pk=c.pk).exists()
    assert Compromiso.all_objects.filter(pk=c.pk).exists()
    operar(c.pk, "restaurar", {}, users["Administrador"])
    assert Compromiso.objects.filter(pk=c.pk).exists()
    operar(c.pk, "eliminar", {}, users["Administrador"])
    with pytest.raises(ValidationError):
        operar(c.pk, "purgar", {"confirmacion": "no"}, users["Administrador"])
    operar(c.pk, "purgar", {"confirmacion": "ELIMINAR"}, users["Administrador"])
    assert not Compromiso.all_objects.filter(pk=c.pk).exists()


@pytest.mark.parametrize("count", [1, 2, 3, 4, 5])
def test_reschedules(c, users, count):
    for i in range(count):
        operar(
            c.pk,
            "reprogramar",
            {"fecha_nueva": date(2026, 9, 11 + i), "motivo": "Dependencia TI"},
            users["Editor"],
        )
    c.refresh_from_db()
    assert c.reprogramaciones.count() == count
    assert fecha_objetivo(c) == date(2026, 9, 10 + min(count, 2))
    assert c.reprogramaciones.last().numero_reprogramacion == count
    assert c.reprogramaciones.last().reprogramacion_clave == f"{c.pk}:{count}"
    assert c.historial.filter(accion="REPROGRAMADO").count() == count


def test_reschedule_rollback(c, users):
    with pytest.raises(ValidationError):
        operar(
            c.pk,
            "reprogramar",
            {"fecha_nueva": date(2026, 8, 1), "motivo": "Adelantar"},
            users["Editor"],
        )
    assert c.reprogramaciones.count() == 0
    assert c.historial.count() == 1


def test_reschedule_legacy(c, users):
    c.segunda_fecha = date(2026, 9, 20)
    c.save()
    operar(
        c.pk,
        "reprogramar",
        {"fecha_nueva": date(2026, 9, 25), "motivo": "Legacy"},
        users["Editor"],
    )
    c.refresh_from_db()
    assert c.tercera_fecha == date(2026, 9, 25)
    assert c.reprogramaciones.get().numero_reprogramacion == 1


@pytest.mark.parametrize(
    "due,real,suspended,expected",
    [
        (date(2026, 9, 10), date(2026, 9, 9), False, "Cumplido a tiempo"),
        (date(2026, 9, 10), date(2026, 9, 12), False, "Cumplido fuera de plazo"),
        (date(2026, 9, 1), None, False, "Vencido"),
        (date(2026, 9, 10), None, False, "Por vencer"),
        (date(2026, 9, 7), None, False, "Por vencer"),
        (date(2026, 9, 14), None, False, "Por vencer"),
        (date(2026, 9, 15), None, False, "En plazo"),
        (None, None, False, "Por definir"),
        (None, date(2026, 9, 8), False, "Por definir"),
        (date(2026, 9, 1), None, True, "Suspendido"),
    ],
)
def test_situation_matches_orm(c, due, real, suspended, expected):
    c.primera_fecha, c.fecha_real, c.suspendida = due, real, suspended
    c.save()
    assert situacion(c, date(2026, 9, 7)) == expected
    assert base(date(2026, 9, 7)).get(pk=c.pk).situacion == expected


def test_target_precedence(c):
    c.primera_fecha, c.segunda_fecha = date(2026, 9, 15), date(2026, 9, 20)
    assert c.fecha_objetivo_vigente == date(2026, 9, 20)


def test_meta(c):
    c.fecha_real = date(2026, 9, 9)
    c.save()
    assert c.cumplimiento_meta == Decimal("105.26")
    assert metricas(base())["cumplimiento_meta"] == pytest.approx(105.26315789)
    c.meta = 0
    c.save()
    assert c.cumplimiento_meta is None
    assert metricas(base())["cumplimiento_meta"] is None


def test_close_suspend_notes(c, users):
    c = operar(
        c.pk,
        "cerrar",
        {"fecha_real": date(2026, 9, 9), "puntaje": 100, "nota": "Listo"},
        users["Editor"],
    )
    assert c.status == settings.CLOSED_STATUS_CODE
    operar(c.pk, "suspender", {"motivo": "En pausa"}, users["Editor"])
    c.refresh_from_db()
    assert c.status == "S" and c.suspendida
    operar(c.pk, "reactivar", {}, users["Editor"])
    operar(c.pk, "nota", {"motivo": "Nueva nota"}, users["Editor"])
    c.refresh_from_db()
    assert not c.suspendida and "Nueva nota" in c.notas


def test_bad_close_atomic(c, users):
    initial_status = c.status
    with pytest.raises(ValidationError):
        operar(
            c.pk,
            "cerrar",
            {"fecha_real": date(2026, 8, 1), "puntaje": 100},
            users["Editor"],
        )
    c.refresh_from_db()
    assert c.status == initial_status and c.fecha_real is None


@pytest.mark.parametrize(
    "params",
    [
        {"q": "formatos"},
        {"tema": "Proyectos"},
        {"iniciativa": "Proyectos Internos"},
        {"responsable_pyp": "HP"},
        {"status": "EC"},
        {"mes": "Septiembre 2026"},
        {"anio": "2026"},
        {"compromiso_hch": "false"},
        {"suspendida": "false"},
        {"desde": "2026-09-10", "hasta": "2026-09-10"},
        {"situacion": "Por vencer"},
    ],
)
def test_filters(c, params):
    assert filtrar(params, date(2026, 9, 7)).count() == 1
    assert filtrar({"q": "inexistente"}).count() == 0


def test_permissions(c, users, client):
    for role in ["Editor", "Consulta"]:
        with pytest.raises(PermissionDenied):
            operar(c.pk, "eliminar", {}, users[role])
    with pytest.raises(PermissionDenied):
        guardar({"tema": "No"}, users["Consulta"], c.pk)
    client.force_login(users["Consulta"])
    for url in [
        "/compromisos/nuevo/",
        f"/compromisos/{c.pk}/editar/",
        f"/compromisos/{c.pk}/cerrar/",
        "/papelera/",
        "/usuarios/",
        "/administracion/",
    ]:
        assert client.get(url).status_code == 403
    assert (
        client.post(f"/compromisos/{c.pk}/suspender/", {"motivo": "x"}).status_code
        == 403
    )


@pytest.mark.parametrize(
    "url",
    [
        "/",
        "/dashboard/",
        "/compromisos/",
        "/indicadores/",
        "/reportes/",
        "/exportar/xlsx/",
        "/exportar/docx/",
    ],
)
def test_read_pages(c, users, client, url):
    assert client.get(url).status_code == 302
    client.force_login(users["Consulta"])
    assert client.get(url).status_code == 200


@pytest.mark.parametrize(
    "url", ["/usuarios/", "/administracion/", "/papelera/", "/compromisos/nuevo/"]
)
def test_admin_pages(c, users, client, url):
    client.force_login(users["Administrador"])
    assert client.get(url).status_code == 200


def test_detail_actions_and_forms(c, users, client, data):
    client.force_login(users["Editor"])
    assert client.get(f"/compromisos/{c.pk}/").status_code == 200
    for action in ["reprogramar", "cerrar", "suspender", "reactivar", "nota"]:
        assert client.get(f"/compromisos/{c.pk}/{action}/").status_code == 200
    assert client.post(f"/compromisos/{c.pk}/reactivar/").status_code == 302
    data["fecha_de_compromiso"] = "2026-09-01"
    data["primera_fecha"] = "2026-09-10"
    assert client.post("/compromisos/nuevo/", data).status_code == 302
    assert (
        client.post(
            f"/compromisos/{c.pk}/editar/", {**data, "tarea": "Editado"}
        ).status_code
        == 302
    )
    c.refresh_from_db()
    assert c.tarea == "Editado"


def test_http_edit_reschedule_and_delete_do_not_conflict_in_history(c, users, client):
    client.force_login(users["Administrador"])
    response = client.post(
        f"/compromisos/{c.pk}/celda/",
        {"field": "tarea", "value": "Edición HTTP", "version": c.updated_at.isoformat()},
        content_type="application/json",
    )
    assert response.status_code == 200

    response = client.post(
        f"/compromisos/{c.pk}/reprogramar/",
        {"fecha_nueva": "2026-09-20", "motivo": "Validación integral"},
    )
    assert response.status_code == 302

    response = client.post(f"/compromisos/{c.pk}/eliminar/", {})
    assert response.status_code == 302
    assert Compromiso.all_objects.get(pk=c.pk).deleted_at is not None


def test_pagination(c, users, client, data):
    for i in range(13):
        guardar({**data, "tarea": f"Tarea {i}"}, users["Editor"])
    client.force_login(users["Consulta"])
    response = client.get("/compromisos/?size=10&tema=Proyectos&page=2")
    assert len(response.context["page_obj"]) == 4
    assert "tema=Proyectos" in response.context["query"]


def test_exports(c):
    c.tarea = '=HYPERLINK("bad")'
    c.save()
    book = load_workbook(BytesIO(excel(base())))
    assert book.active.max_column == 17
    assert book.active["D2"].data_type == "s"
    doc = Document(BytesIO(word(base(), date(2026, 9, 7))))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "07/09/2026" in text and "Proyectos Internos" in text


def test_idempotent_seed(users):
    call_command("init_portal")
    call_command("seed_demo")
    n = Compromiso.objects.count()
    call_command("seed_demo")
    assert Compromiso.objects.count() == n
    assert (
        Group.objects.filter(name__in=["Administrador", "Editor", "Consulta"]).count()
        == 3
    )


def test_metrics_denominators(c, users, data):
    operar(
        c.pk,
        "cerrar",
        {"fecha_real": date(2026, 9, 9), "puntaje": 100},
        users["Editor"],
    )
    guardar({**data, "suspendida": True}, users["Editor"])
    k = metricas(base(date(2026, 9, 7)))
    assert k["cumplimiento_cierre"] == 100
    assert k["cumplimiento_plazo"] == 100
    assert k["total"] == 2


def test_csrf(c, users):
    from django.test import Client

    client = Client(enforce_csrf_checks=True)
    client.force_login(users["Editor"])
    assert client.post(f"/compromisos/{c.pk}/reactivar/").status_code == 403
