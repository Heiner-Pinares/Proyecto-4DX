from datetime import date
from io import BytesIO

import pytest
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.test import Client
from openpyxl import load_workbook

from compromisos.indicators import metricas
from compromisos.models import Compromiso
from compromisos.reports import excel
from compromisos.scoring import calcular_puntaje
from compromisos.selectors import base

pytestmark = pytest.mark.django_db


@pytest.fixture
def editor():
    call_command("init_portal")
    u = User.objects.create_user("grid_editor", password="test-password")
    u.groups.add(Group.objects.get(name="Editor"))
    return u


@pytest.fixture
def c():
    return Compromiso.objects.create(
        tema="Proyectos",
        iniciativa="Portal",
        tarea="Prueba de edición",
        responsable_pyp="HP",
        status="EC",
        fecha_de_compromiso=date(2026, 9, 1),
        primera_fecha=date(2026, 9, 10),
        fecha_de_vencimiento=date(2026, 9, 10),
    )


@pytest.mark.parametrize(
    "real,second,third,expected",
    [
        (9, None, None, 100),
        (10, None, None, 100),
        (11, None, None, 60),
        (12, 15, None, 80),
        (15, 15, None, 80),
        (16, 15, None, 60),
        (18, 15, 20, 60),
        (20, 15, 20, 60),
        (25, 15, 20, 60),
        (9, 15, 20, 100),
        (12, 15, 20, 80),
        (18, None, 20, 60),
        (None, 15, 20, None),
    ],
)
def test_score_rules(c, real, second, third, expected):
    c.fecha_real = date(2026, 9, real) if real else None
    c.segunda_fecha = date(2026, 9, second) if second else None
    c.tercera_fecha = date(2026, 9, third) if third else None
    c.save()
    c.refresh_from_db()
    assert c.puntaje == expected
    assert calcular_puntaje(c) == expected
    assert base().get(pk=c.pk).cumplimiento_fechas == expected
    assert load_workbook(BytesIO(excel(base()))).active["L2"].value == expected


def test_initial_date_and_no_target(c):
    c.primera_fecha = date(2026, 9, 10)
    c.fecha_real = date(2026, 9, 10)
    c.save()
    assert c.puntaje == 100
    c.primera_fecha = None
    c.save()
    assert c.puntaje == 60


def post(client, c, field, value, **extra):
    return client.post(
        f"/compromisos/{c.pk}/celda/",
        {"field": field, "value": value, "version": c.updated_at.isoformat(), **extra},
        content_type="application/json",
    )


def test_save_recalculate_conflict(c, editor, client):
    client.force_login(editor)
    response = post(client, c, "fecha_real", "2026-09-11")
    assert response.status_code == 200
    assert response.json()["row"]["cells"][11]["display"] == "60 %"
    assert post(client, c, "tarea", "Se perdió esta edición").status_code == 409
    c.refresh_from_db()
    assert c.puntaje == 60
    response = post(
        client, c, "segunda_fecha", "2026-09-15", reason="Acuerdo del equipo"
    )
    assert response.status_code == 200
    c.refresh_from_db()
    assert c.puntaje == 80
    assert c.reprogramaciones.count() == 1
    assert "60" in c.historial.first().descripcion
    assert "80" in c.historial.first().descripcion


@pytest.mark.parametrize(
    "field,value",
    [
        ("tema", ""),
        ("meta", -1),
        ("status", "FALSO"),
        ("puntaje", 100),
        ("mes", "Otro"),
        ("created_by", "atacante"),
        ("fecha_real", "2026-08-01"),
        ("segunda_fecha", "2026-09-05"),
    ],
)
def test_validation(c, editor, client, field, value):
    client.force_login(editor)
    assert post(client, c, field, value, reason="Prueba").status_code == 400
    assert c.historial.count() == 0


def test_reason_and_permissions(c, editor, client):
    client.force_login(editor)
    assert post(client, c, "segunda_fecha", "2026-09-15").status_code == 400
    editor.groups.clear()
    editor.groups.add(Group.objects.get(name="Consulta"))
    assert post(client, c, "tarea", "Prohibido").status_code == 403
    html = client.get("/compromisos/").content.decode()
    assert 'class="cell-value"' not in html


def test_csrf(c, editor):
    client = Client(enforce_csrf_checks=True)
    client.force_login(editor)
    assert post(client, c, "tarea", "Sin CSRF").status_code == 403


def test_automatic_month_and_readonly_score(c, editor, client):
    client.force_login(editor)
    assert post(client, c, "fecha_de_compromiso", "2026-08-01").status_code == 200
    c.refresh_from_db()
    assert c.mes == "Agosto 2026"
    assert client.get(f"/compromisos/{c.pk}/celda/").status_code == 405
    assert (
        client.post(
            f"/compromisos/{c.pk}/celda/", "[1]", content_type="application/json"
        ).status_code
        == 400
    )


def test_average_excludes_open_and_suspended(c):
    c.fecha_real = date(2026, 9, 11)
    c.save()
    Compromiso.objects.create(
        tema="A", iniciativa="B", tarea="Sin cerrar", responsable_pyp="HP", status="EC"
    )
    Compromiso.objects.create(
        tema="A",
        iniciativa="B",
        tarea="Suspendido",
        responsable_pyp="HP",
        status="T",
        suspendida=True,
        fecha_real=date(2026, 9, 1),
        primera_fecha=date(2026, 9, 1),
    )
    assert metricas(base())["cumplimiento_fechas"] == 60


def test_automatic_deadline_and_score(c, client, editor):
    client.force_login(editor)
    assert post(client, c, 'fecha_de_vencimiento', '2026-10-01').status_code == 400
    result = post(client, c, 'segunda_fecha', '2026-09-15', reason='Dependencia')
    assert result.status_code == 200
    c.refresh_from_db()
    assert c.fecha_de_vencimiento == date(2026, 9, 15)
    c.fecha_real = date(2026, 9, 15)
    c.save()
    assert c.puntaje == 80
    assert base().get(pk=c.pk).cumplimiento_fechas == 80
    assert post(client, c, 'tercera_fecha', '2026-09-20', reason='Ampliación').status_code == 200
    c.refresh_from_db()
    assert c.fecha_de_vencimiento == date(2026, 9, 20)
    assert post(client, c, 'tercera_fecha', '', reason='Corrección').status_code == 200
    c.refresh_from_db()
    assert c.fecha_de_vencimiento == date(2026, 9, 15)


def test_status_and_suspension_sync(c, client, editor):
    client.force_login(editor)
    for status in ['S', 'EC', 'D', 'T']:
        assert post(client, c, 'status', status).status_code == 400
    assert post(client, c, 'fecha_real', '2026-09-10').status_code == 200
    c.refresh_from_db()
    assert c.status == 'T'
    assert post(client, c, 'suspendida', True).status_code == 200
    c.refresh_from_db()
    assert c.status == 'S'
    assert post(client, c, 'suspendida', False).status_code == 200
    c.refresh_from_db()
    assert c.status == 'T'


@pytest.mark.parametrize('first,second,third,real,suspended,expected', [
    (None,None,None,None,False,'EC'),
    (9,None,None,None,False,'D'),
    (10,None,None,None,False,'EC'),
    (9,11,None,None,False,'EC'),
    (8,9,None,None,False,'D'),
    (8,9,11,None,False,'EC'),
    (7,8,9,None,False,'D'),
    (7,8,9,10,False,'T'),
    (7,8,9,10,True,'S'),
    (7,8,9,None,True,'S'),
])
def test_automatic_status_rules(c, first, second, third, real, suspended, expected):
    from compromisos.status import calcular_estatus
    c.primera_fecha = date(2026,9,first) if first else None
    c.segunda_fecha = date(2026,9,second) if second else None
    c.tercera_fecha = date(2026,9,third) if third else None
    c.fecha_real = date(2026,9,real) if real else None
    c.suspendida = suspended
    c.save()
    assert calcular_estatus(c, date(2026,9,10)) == expected
    assert base(date(2026,9,10)).get(pk=c.pk).estatus == expected


def test_status_advances_without_editing(c):
    from compromisos.selectors import filtrar
    from compromisos.grid import fila
    from compromisos.models import Compromiso
    Compromiso.objects.filter(pk=c.pk).update(status='EC')
    before = base(date(2026,9,10)).get(pk=c.pk)
    after = base(date(2026,9,11)).get(pk=c.pk)
    assert before.estatus == 'EC' and after.estatus == 'D'
    assert filtrar({'status':'D'}, date(2026,9,11)).filter(pk=c.pk).exists()
    assert next(cell for cell in fila(after)['cells'] if cell['name']=='status')['value'] == 'D'
    assert load_workbook(BytesIO(excel(base(date(2026,9,11))))).active['E2'].value == 'D'
