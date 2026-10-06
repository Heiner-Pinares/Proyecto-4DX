from datetime import date
from io import BytesIO
from itertools import combinations

import pytest
from django.contrib.auth.models import User
from django.http import QueryDict
from openpyxl import load_workbook

from compromisos.models import Compromiso
from compromisos.selectors import filtrar

pytestmark = pytest.mark.django_db


@pytest.fixture
def etapas():
    result = {}
    for n in range(4):
        result[str(n)] = Compromiso.objects.create(
            proyecto='Proyecto', tarea=f'Etapa {n}', responsable_pyp='HP',
            primera_fecha=date(2026,9,1) if n >= 1 else None,
            segunda_fecha=date(2026,9,10) if n >= 2 else None,
            tercera_fecha=date(2026,9,20) if n >= 3 else None,
        )
    return result


@pytest.mark.parametrize('selected', [list(c) for n in range(4) for c in combinations('123', n)])
def test_multi_select_is_union_of_latest_stage(etapas, selected):
    params = QueryDict('', mutable=True)
    params.setlist('etapa_fecha', selected)
    expected = selected or list(etapas)
    assert set(filtrar(params).values_list('pk', flat=True)) == {etapas[x].pk for x in expected}


def test_render_pagination_export_and_other_filters(client, etapas):
    client.force_login(User.objects.create_user('consulta_fechas'))
    query = 'etapa_fecha=1&etapa_fecha=3'
    response = client.get('/compromisos/?'+query)
    assert response.status_code == 200
    assert response.context['etapas_fecha'] == ['1','3']
    assert response.context['query'].count('etapa_fecha=') == 2
    html = response.content.decode()
    assert 'value="1" checked' in html and 'value="3" checked' in html
    assert {r['id'] for r in response.context['rows']} == {etapas['1'].pk, etapas['3'].pk}
    book = load_workbook(BytesIO(client.get('/exportar/xlsx/?'+query).content))
    assert book.active.max_row == 3
    assert {book.active.cell(n,2).value for n in (2,3)} == {'Etapa 1','Etapa 3'}
    assert filtrar({'etapa_fecha':['1','3'], 'q':'Etapa 3'}).get().pk == etapas['3'].pk
    assert filtrar({'etapa_fecha':'2'}).get().pk == etapas['2'].pk
