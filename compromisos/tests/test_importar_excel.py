from datetime import datetime

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import Workbook, load_workbook

from compromisos.models import Compromiso, EventoCompromiso


HEADERS = [
    "Nro Proy.", "Proyecto", "Estado Proyecto", "ID Actividad", "Actividad",
    "Jefatura ", "Responsable", "Fecha original", "1ra reprogramación",
    "2da reprogramación", "Fecha compromiso (FET)", "Reprogramada",
    "Motivo de reprogramación", "Estado", "Cumplimiento", "Situación",
]


def workbook(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Hoja1"
    sheet.append(HEADERS)
    sheet.append([
        5, "Recibo PDF", "En curso", "5.1", "Publicar recibo",
        "Soporte Post Facturación", "Ana", datetime(2026, 9, 4),
        datetime(2026, 9, 11), None, None, None, "Ajuste de TI", "Cerrado",
        0.8, "Cerrado",
    ])
    sheet.append([
        25, "Automatización", "En curso", "5.1", "Completar pruebas",
        None, None, None, None, None, None, None, None, "Por definir", None,
        "Sin fecha",
    ])
    book.save(path)


@pytest.mark.django_db
def test_import_is_dry_run_by_default_and_idempotent(tmp_path):
    path = tmp_path / "compromisos.xlsx"
    workbook(path)

    call_command("importar_compromisos_excel", str(path), verbosity=0)
    assert Compromiso.objects.count() == 0

    call_command("importar_compromisos_excel", str(path), aplicar=True, verbosity=0)
    assert Compromiso.objects.count() == 2
    closed = Compromiso.objects.get(codigo_fuente="excel4dx:5:5.1")
    assert closed.proyecto == "Recibo PDF"
    assert closed.iniciativa == ""
    assert closed.jefatura == "Jefatura de Soporte Oper Post Facturacion"
    assert closed.status == "T"
    assert closed.fecha_real.isoformat() == "2026-09-11"
    assert closed.puntaje == 80
    assert closed.reprogramaciones.get().fecha_anterior.isoformat() == "2026-09-04"
    unassigned = Compromiso.objects.get(codigo_fuente="excel4dx:25:5.1")
    assert unassigned.responsable_pyp == "Por asignar"
    assert unassigned.jefatura == ""
    assert unassigned.fecha_de_vencimiento is None
    assert EventoCompromiso.objects.filter(accion="IMPORTADO").count() == 2

    call_command("importar_compromisos_excel", str(path), aplicar=True, verbosity=0)
    assert Compromiso.objects.count() == 2
    assert EventoCompromiso.objects.filter(accion="IMPORTADO").count() == 2


@pytest.mark.django_db
def test_replace_all_validates_previews_and_requires_confirmation(tmp_path):
    path = tmp_path / "compromisos.xlsx"
    workbook(path)
    call_command("importar_compromisos_excel", str(path), aplicar=True, verbosity=0)
    assert Compromiso.objects.count() == 2

    book = load_workbook(path)
    book["Hoja1"].delete_rows(3)
    book.save(path)

    call_command(
        "importar_compromisos_excel",
        str(path),
        reemplazar_todo=True,
        verbosity=0,
    )
    assert Compromiso.objects.count() == 2

    with pytest.raises(CommandError, match="--confirmar REEMPLAZAR"):
        call_command(
            "importar_compromisos_excel",
            str(path),
            aplicar=True,
            reemplazar_todo=True,
            verbosity=0,
        )
    assert Compromiso.objects.count() == 2

    call_command(
        "importar_compromisos_excel",
        str(path),
        aplicar=True,
        reemplazar_todo=True,
        confirmar="REEMPLAZAR",
        verbosity=0,
    )
    assert list(Compromiso.objects.values_list("codigo_fuente", flat=True)) == [
        "excel4dx:5:5.1"
    ]
    assert EventoCompromiso.objects.filter(
        tipo="historial", accion="IMPORTADO"
    ).count() == 1
