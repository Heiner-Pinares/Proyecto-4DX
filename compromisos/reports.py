"""Exportaciones generadas en memoria, con nombres fijos y texto seguro para Excel."""

import logging
from io import BytesIO
from itertools import groupby

from django.conf import settings
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .services import situacion

log = logging.getLogger(__name__)
LEGACY = [
    ("Compromisos de HCH", "compromiso_hch"),
    ("TEMA", "tema"),
    ("INICIATIVA", "iniciativa"),
    ("TAREA", "tarea"),
    ("STATUS", "status"),
    ("Meta", "meta"),
    ("Responsable_PyP", "responsable_pyp"),
    ("Mes", "mes"),
    ("FECHA_DE_COMPROMISO", "fecha_de_compromiso"),
    ("FECHA_DE_VENCIMIENTO", "fecha_de_vencimiento"),
    ("FECHA REAL", "fecha_real"),
    ("PUNTAJE", "puntaje"),
    ("1ERA FECHA", "primera_fecha"),
    ("2DA FECHA", "segunda_fecha"),
    ("3RA FECHA", "tercera_fecha"),
    ("NOTAS", "notas"),
    ("Suspendida", "suspendida"),
]


def excel(qs):
    book = Workbook()
    sheet = book.active
    sheet.title = "Compromisos"
    sheet.append([x[0] for x in LEGACY])
    for c in qs.iterator(chunk_size=500):
        sheet.append([c.estatus_actual if field == "status" else getattr(c, field) for _, field in LEGACY])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = (
                    "s"  # Evita fórmulas introducidas mediante texto del usuario.
                )
            if hasattr(cell.value, "year"):
                cell.number_format = "DD/MM/YYYY"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="000000")
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = (
            24 if col[0].column != 4 else 70
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    log.info("Exportación Excel generada")
    return output.getvalue()


def word(qs, corte):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Reporte ejecutivo de compromisos", 0)
    doc.add_paragraph(f"Estados actualizados al cierre del {corte:%d/%m/%Y}")
    doc.add_paragraph(
        "Seguimiento agrupado por tema e iniciativa. La situación del plazo se evalúa a la fecha de corte sobre los datos actuales; este reporte no reconstruye versiones históricas."
    )
    rows = qs.order_by("tema", "iniciativa", "id")

    def fmt(d):
        return d.strftime("%d/%m/%Y") if d else "Por definir"

    for tema, temas in groupby(rows, key=lambda c: c.tema):
        doc.add_heading(tema, 1)
        for iniciativa, items in groupby(temas, key=lambda c: c.iniciativa):
            doc.add_heading(iniciativa, 2)
            for c in items:
                cierre = (
                    "Cerrado" if c.estatus_actual == settings.CLOSED_STATUS_CODE else c.estatus_actual
                )
                p = doc.add_paragraph()
                p.add_run(c.tarea).bold = True
                doc.add_paragraph(
                    f"Responsable: {c.responsable_pyp} · Estatus: {cierre} · {situacion(c, corte)}"
                )
                doc.add_paragraph(
                    f"Vencimiento vigente: {fmt(c.fecha_de_vencimiento)} · Objetivo vigente: {fmt(c.fecha_objetivo_vigente)} · Fecha real: {fmt(c.fecha_real)}"
                )
                doc.add_paragraph(
                    f"Cumplimiento de fechas: {str(c.puntaje) + ' %' if c.puntaje is not None else 'Pendiente de cierre'}"
                )
                history = list(c.reprogramaciones.all())
                for n, d in enumerate(
                    [c.primera_fecha, c.segunda_fecha, c.tercera_fecha], 1
                ):
                    if d and not any(r.numero_reprogramacion == n for r in history):
                        doc.add_paragraph(
                            f"Reprogramación {n}: {fmt(d)} (registro legacy)",
                            style="List Bullet",
                        )
                for r in history:
                    doc.add_paragraph(
                        f"Reprogramación {r.numero_reprogramacion}: {fmt(r.fecha_anterior)} → {fmt(r.fecha_nueva)}. {r.motivo}",
                        style="List Bullet",
                    )
                if c.notas:
                    doc.add_paragraph("Observaciones: " + c.notas)
    if not rows:
        doc.add_paragraph("No hay compromisos para los filtros seleccionados.")
    output = BytesIO()
    doc.save(output)
    log.info("Reporte ejecutivo generado con corte %s", corte)
    return output.getvalue()


def correo_contexto(qs, corte):
    """Pendientes actuales iniciados hasta el corte; no reconstruye historia."""
    from django.db.models import Q

    rows = list(qs.filter(suspendida=False, fecha_real__isnull=True)
                .exclude(estatus=settings.CLOSED_STATUS_CODE)
                .filter(Q(fecha_de_compromiso__lte=corte) | Q(fecha_de_compromiso__isnull=True))
                .order_by("tema", "iniciativa", "objetivo", "id"))
    resumen = {"total": len(rows), "vencidos": 0, "por_vencer": 0, "en_plazo": 0, "sin_fecha": 0}
    colors = {"Vencido": "#DA291C", "Por vencer": "#FFCB00", "En plazo": "#009FDF", "Por definir": "#FF67B9"}
    keys = {"Vencido": "vencidos", "Por vencer": "por_vencer", "En plazo": "en_plazo", "Por definir": "sin_fecha"}
    for row in rows:
        resumen[keys[row.situacion]] += 1
        row.correo_color = colors[row.situacion]
        row.correo_texto = "#FFFFFF" if row.situacion == "Vencido" else "#000000"
        row.dias_atraso = (corte - row.objetivo).days if row.objetivo and row.objetivo < corte else 0
    return {"rows": rows, "resumen": resumen, "corte": corte,
            "asunto": f"Seguimiento de compromisos en proceso | Corte {corte:%d/%m/%Y}"}


def resumen_jefaturas(qs, corte):
    from datetime import timedelta
    from django.db.models import Q
    from .models import Compromiso
    from .scoring import calcular_puntaje

    fin = corte + timedelta(days=6)
    rows = list(qs.filter(suspendida=False).filter(
        Q(fecha_de_compromiso__lte=corte) | Q(fecha_de_compromiso__isnull=True)
    ).order_by('objetivo', 'id'))
    palette = [('#DA291C','#FFF0F1'),('#996000','#FFF8DF'),('#007AA8','#EAF8FF')]
    groups = {name: {'nombre': name, 'color': palette[i][0], 'fondo': palette[i][1], 'proximos': [], 'scores': []}
              for i, (name, _) in enumerate(Compromiso.JEFATURAS)}
    scores = []
    for row in rows:
        name = row.jefatura or 'Sin jefatura asignada'
        group = groups.setdefault(name, {'nombre':name,'color':'#000000','fondo':'#F4F4F4','proximos':[],'scores':[]})
        if not row.fecha_real and row.objetivo and corte <= row.objetivo <= fin:
            group['proximos'].append(row)
        if row.fecha_real and row.fecha_real <= corte:
            score = calcular_puntaje(row)
            group['scores'].append(score)
            scores.append(score)
    for group in groups.values():
        group['cumplimiento'] = sum(group['scores']) / len(group['scores']) if group['scores'] else None
        group['cerrados'] = len(group['scores'])
    return {'grupos':list(groups.values()), 'corte':corte, 'fin':fin,
            'general':sum(scores)/len(scores) if scores else None,
            'cerrados':len(scores), 'total_proximos':sum(len(g['proximos']) for g in groups.values())}
