"""Crea y actualiza cuentas del portal a partir del padrón de personal."""

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata

from django.contrib.auth.models import Group, User
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook

from compromisos.profiles import save_profile


HEADERS = {
    "cod comunicacion": "code",
    "correo electronico": "email",
    "jefe": "boss",
    "nombre completo": "full_name",
    "area": "area",
    "gerencia": "management",
    "direccion": "division",
}


def normalized(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def username_piece(value):
    return re.sub(r"[^a-z0-9]+", "", normalized(value))


def first_surname(full_name, email):
    """Usa el alias corporativo para reconocer apellidos compuestos."""
    parts = full_name.split()
    raw_local = unicodedata.normalize("NFKD", email.split("@", 1)[0])
    raw_local = "".join(char for char in raw_local if not unicodedata.combining(char)).lower()
    segments = [value for value in re.split(r"[^a-z0-9]+", raw_local) if value]
    compact_local = "".join(segments)
    first = username_piece(parts[0])
    if len(segments) > 1:
        alias_surname = segments[-1]
    elif compact_local.startswith(first) and len(compact_local) > len(first):
        alias_surname = compact_local[len(first) :]
    elif compact_local.startswith(first[:1]) and len(compact_local) > 1:
        alias_surname = compact_local[1:]
    else:
        alias_surname = ""
    candidates = []
    for start in range(1, len(parts)):
        for end in range(start + 1, len(parts) + 1):
            candidate = " ".join(parts[start:end])
            compact = username_piece(candidate)
            if len(compact) >= 3 and alias_surname.startswith(compact):
                candidates.append((len(compact), -start, candidate))
    if candidates:
        return max(candidates)[2]
    # Respaldo para alias numéricos: apellido anterior al materno, incluyendo
    # partículas frecuentes como "Del Castillo".
    index = len(parts) - 2
    start = index
    while start > 1 and normalized(parts[start - 1]) in {"de", "del", "la", "las", "los"}:
        start -= 1
    return " ".join(parts[start : index + 1])


@dataclass(frozen=True)
class Person:
    row: int
    code: str
    email: str
    full_name: str
    first_name: str
    surname: str
    boss: str
    area: str
    management: str
    division: str
    base_username: str
    username: str = ""


def read_people(path, sheet_name):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise CommandError(f"No se pudo abrir el Excel: {error}") from error
    try:
        if sheet_name not in workbook.sheetnames:
            raise CommandError(
                f"No existe la hoja {sheet_name!r}. Disponibles: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        positions = {
            HEADERS[normalized(value)]: index
            for index, value in enumerate(header)
            if normalized(value) in HEADERS
        }
        missing = sorted(set(HEADERS.values()) - set(positions))
        if missing:
            raise CommandError("Faltan columnas requeridas: " + ", ".join(missing))

        people = []
        errors = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            get = lambda name: clean(values[positions[name]])
            if not any(get(name) for name in positions):
                continue
            code = get("code")
            email = get("email").lower()
            full_name = get("full_name")
            if not code:
                errors.append(f"fila {row_number}: falta Cod. Comunicación")
                continue
            try:
                validate_email(email)
            except ValidationError:
                errors.append(f"fila {row_number}: correo electrónico no válido")
                continue
            parts = full_name.split()
            if len(parts) < 3:
                errors.append(f"fila {row_number}: Nombre completo no permite identificar nombre y apellido")
                continue
            surname = first_surname(full_name, email)
            first_name = parts[0]
            base_username = f"{username_piece(first_name)}.{username_piece(surname)}"
            if not base_username.strip("."):
                errors.append(f"fila {row_number}: no se pudo formar el usuario")
                continue
            people.append(
                Person(
                    row=row_number,
                    code=code,
                    email=email,
                    full_name=full_name,
                    first_name=first_name.title(),
                    surname=surname.title(),
                    boss=get("boss"),
                    area=get("area"),
                    management=get("management"),
                    division=get("division"),
                    base_username=base_username,
                )
            )
    finally:
        workbook.close()

    for label, values in [
        ("Cod. Comunicación", [p.code.casefold() for p in people]),
        ("Correo electrónico", [p.email.casefold() for p in people]),
        ("Nombre completo", [normalized(p.full_name) for p in people]),
    ]:
        duplicates = [value for value, total in Counter(values).items() if total > 1]
        if duplicates:
            errors.append(f"{label} contiene {len(duplicates)} valor(es) repetido(s)")
    if errors:
        raise CommandError("El Excel contiene errores:\n- " + "\n- ".join(errors))

    groups = defaultdict(list)
    for person in people:
        groups[person.base_username].append(person)
    assigned = []
    for base, members in groups.items():
        for position, person in enumerate(sorted(members, key=lambda item: item.email), start=1):
            username = base if position == 1 else f"{base}{position}"
            assigned.append(Person(**{**person.__dict__, "username": username}))
    return sorted(assigned, key=lambda person: person.row)


class Command(BaseCommand):
    help = "Importa usuarios, áreas y estructura organizacional; las contraseñas se guardan con hash."

    def add_arguments(self, parser):
        parser.add_argument("archivo", help="Ruta del archivo Operaciones Comerciales.xlsx")
        parser.add_argument("--hoja", default="OPERACIONES")
        parser.add_argument("--rol", default="Consulta", choices=["Administrador", "Editor", "Consulta"])
        parser.add_argument("--solo-validar", action="store_true")
        parser.add_argument("--restablecer-contrasenas", action="store_true")
        parser.add_argument("--actualizar-rol", action="store_true")

    def handle(self, *args, **options):
        path = Path(options["archivo"]).expanduser().resolve()
        if not path.is_file():
            raise CommandError(f"No existe el archivo: {path}")
        people = read_people(path, options["hoja"])
        collisions = len(people) - len({p.base_username for p in people})
        if options["solo_validar"]:
            self.stdout.write(
                self.style.SUCCESS(
                    f"Archivo válido: {len(people)} personas, {collisions} colisión(es) de usuario resuelta(s)."
                )
            )
            return
        try:
            role = Group.objects.get(name=options["rol"])
        except Group.DoesNotExist as error:
            raise CommandError("Faltan los roles del portal. Ejecuta primero: python run.py --check") from error

        created = updated = password_changes = 0
        with transaction.atomic():
            for person in people:
                by_email = User.objects.filter(email__iexact=person.email).first()
                by_username = User.objects.filter(username=person.username).first()
                if by_email and by_username and by_email.pk != by_username.pk:
                    raise CommandError(
                        f"fila {person.row}: el correo y el usuario pertenecen a cuentas diferentes"
                    )
                user = by_email or by_username
                is_new = user is None
                if is_new:
                    user = User(username=person.username)
                elif user.email and user.email.casefold() != person.email.casefold():
                    raise CommandError(
                        f"fila {person.row}: el usuario calculado ya existe con otro correo"
                    )
                user.username = person.username
                user.first_name = person.first_name
                user.last_name = person.surname
                user.email = person.email
                user.is_active = True
                if is_new or options["restablecer_contrasenas"]:
                    user.set_password(person.code)
                    password_changes += 1
                user.save()
                if is_new or options["actualizar_rol"] or not user.groups.exists():
                    user.groups.set([role])
                save_profile(
                    user,
                    nombre_completo=person.full_name,
                    area=person.area,
                    gerencia=person.management,
                    direccion=person.division,
                    jefe=person.boss,
                )
                if is_new:
                    created += 1
                else:
                    updated += 1
        self.stdout.write(
            self.style.SUCCESS(
                f"Importación completa: {created} creados, {updated} actualizados, "
                f"{password_changes} contraseñas establecidas y {collisions} colisión(es) resuelta(s)."
            )
        )
