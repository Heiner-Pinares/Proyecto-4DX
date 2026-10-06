"""Importa la matriz oficial de actividades 4DX sin duplicar registros."""

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
import unicodedata

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from compromisos.models import Compromiso, EventoCompromiso


JEFATURAS = {
    "facturacion a clientes": "Jefatura de Facturacion a Clientes",
    "jefatura de facturacion a clientes": "Jefatura de Facturacion a Clientes",
    "soporte post facturacion": "Jefatura de Soporte Oper Post Facturacion",
    "soporte oper post facturacion": "Jefatura de Soporte Oper Post Facturacion",
    "jefatura de soporte oper post facturacion": "Jefatura de Soporte Oper Post Facturacion",
    "control y mejoras": "Jefatura de Control y Mejoras de Facturacion",
    "control y mejoras de facturacion": "Jefatura de Control y Mejoras de Facturacion",
    "jefatura de control y mejoras de facturacion": "Jefatura de Control y Mejoras de Facturacion",
}

REQUIRED_HEADERS = {
    "nro proy": "numero_proyecto",
    "proyecto": "proyecto",
    "estado proyecto": "estado_proyecto",
    "id actividad": "id_actividad",
    "actividad": "actividad",
    "jefatura": "jefatura",
    "responsable": "responsable",
    "fecha original": "primera_fecha",
    "1ra reprogramacion": "segunda_fecha",
    "2da reprogramacion": "tercera_fecha",
    "motivo de reprogramacion": "motivo",
    "estado": "estado",
}


def normalized(value):
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def clean_text(value):
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def clean_identifier(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def excel_date(value, row, header):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(value, pattern).date()
            except ValueError:
                pass
    raise ValueError(f"fila {row}: {header} no contiene una fecha válida")


@dataclass(frozen=True)
class ExcelCommitment:
    row: int
    source_key: str
    project_number: str
    activity_id: str
    project: str
    project_status: str
    activity: str
    department: str
    owner: str
    first_date: date | None
    second_date: date | None
    third_date: date | None
    reason: str
    source_status: str

    @property
    def current_date(self):
        return self.third_date or self.second_date or self.first_date


def read_rows(path, sheet_name="Hoja1"):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise CommandError(f"No se pudo abrir el Excel: {error}") from error
    if sheet_name not in workbook.sheetnames:
        raise CommandError(
            f"No existe la hoja {sheet_name!r}. Disponibles: {', '.join(workbook.sheetnames)}"
        )
    sheet = workbook[sheet_name]
    header_positions = {}
    for column, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))):
        name = normalized(cell).replace("nro proy ", "nro proy")
        if name in REQUIRED_HEADERS:
            header_positions[REQUIRED_HEADERS[name]] = column
    missing = sorted(set(REQUIRED_HEADERS.values()) - set(header_positions))
    if missing:
        raise CommandError("Faltan columnas requeridas: " + ", ".join(missing))

    parsed = []
    errors = []
    warnings = {"sin_jefatura": [], "sin_responsable": [], "sin_fecha": [], "proyecto_sin_cache": []}
    keys = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda name: values[header_positions[name]]
        activity = clean_text(get("actividad"))
        if not activity:
            continue
        project_number = clean_identifier(get("numero_proyecto"))
        activity_id = clean_identifier(get("id_actividad"))
        if not project_number or not activity_id:
            errors.append(f"fila {row_number}: falta Nro Proy. o ID Actividad")
            continue
        source_key = f"excel4dx:{project_number}:{activity_id}"
        if source_key in keys:
            errors.append(f"fila {row_number}: clave repetida {source_key}")
            continue
        keys.add(source_key)

        project = clean_text(get("proyecto"))
        if not project:
            project = f"Proyecto {project_number}"
            warnings["proyecto_sin_cache"].append(row_number)
        raw_department = clean_text(get("jefatura"))
        department = JEFATURAS.get(normalized(raw_department), "")
        if raw_department and not department:
            errors.append(f"fila {row_number}: jefatura desconocida {raw_department!r}")
            continue
        if not department:
            warnings["sin_jefatura"].append(row_number)
        owner = clean_text(get("responsable"))
        if not owner:
            owner = "Por asignar"
            warnings["sin_responsable"].append(row_number)
        try:
            first = excel_date(get("primera_fecha"), row_number, "Fecha original")
            second = excel_date(get("segunda_fecha"), row_number, "1ra reprogramación")
            third = excel_date(get("tercera_fecha"), row_number, "2da reprogramación")
        except ValueError as error:
            errors.append(str(error))
            continue
        if second and (not first or second < first):
            errors.append(f"fila {row_number}: la 1ra reprogramación precede a la fecha original")
            continue
        if third and (not second or third < second):
            errors.append(f"fila {row_number}: la 2da reprogramación no sigue a la primera")
            continue
        if not first:
            warnings["sin_fecha"].append(row_number)
        source_status = normalized(get("estado"))
        if source_status not in {"cerrado", "pendiente", "por definir", "stand by", "suspendido"}:
            errors.append(f"fila {row_number}: estado desconocido {clean_text(get('estado'))!r}")
            continue
        current = third or second or first
        if source_status == "cerrado" and not current:
            errors.append(
                f"fila {row_number}: una actividad cerrada necesita al menos una fecha comprometida"
            )
            continue
        parsed.append(
            ExcelCommitment(
                row=row_number,
                source_key=source_key,
                project_number=project_number,
                activity_id=activity_id,
                project=project,
                project_status=clean_text(get("estado_proyecto")),
                activity=activity,
                department=department,
                owner=owner,
                first_date=first,
                second_date=second,
                third_date=third,
                reason=clean_text(get("motivo")),
                source_status=source_status,
            )
        )
    workbook.close()
    if errors:
        raise CommandError("El Excel contiene errores:\n- " + "\n- ".join(errors))
    return parsed, warnings


def fields_for(item, filename, sheet_name):
    closed = item.source_status == "cerrado"
    suspended = item.source_status in {"stand by", "suspendido"}
    notes = [
        f"Origen Excel: {filename}; hoja {sheet_name}; fila {item.row}.",
        f"Proyecto {item.project_number}; actividad {item.activity_id}.",
        f"Estado del proyecto en la fuente: {item.project_status or 'Sin dato'}.",
        f"Estado de la actividad en la fuente: {item.source_status.title()}.",
    ]
    if item.reason:
        notes.append(f"Motivo de reprogramación: {item.reason}.")
    if closed:
        notes.append(
            "La fecha real se tomó de la última fecha comprometida porque el Excel no incluye una fecha de cierre separada."
        )
    return {
        "codigo_fuente": item.source_key,
        "proyecto": item.project,
        "iniciativa": "",
        "tarea": item.activity,
        "status": "EC",
        "responsable_pyp": item.owner,
        "jefatura": item.department,
        "fecha_de_compromiso": item.first_date,
        "primera_fecha": item.first_date,
        "segunda_fecha": item.second_date,
        "tercera_fecha": item.third_date,
        "fecha_real": item.current_date if closed else None,
        "suspendida": suspended,
        "notas": "\n".join(notes),
        "created_by": "importacion_excel",
        "updated_by": "importacion_excel",
    }


def save_events(commitment, item, filename, sheet_name):
    dates = [(1, item.first_date, item.second_date), (2, item.second_date, item.third_date)]
    for number, previous, new in dates:
        if not new:
            continue
        EventoCompromiso.objects.update_or_create(
            compromiso=commitment,
            numero_reprogramacion=number,
            defaults={
                "tipo": "reprogramacion",
                "fecha_anterior": previous,
                "fecha_nueva": new,
                "motivo": item.reason or "Sin motivo registrado en el Excel de origen.",
                "usuario": "importacion_excel",
            },
        )
    imported = EventoCompromiso.objects.filter(
        compromiso=commitment,
        tipo="historial",
        accion="IMPORTADO",
        usuario="importacion_excel",
    ).exists()
    if not imported:
        EventoCompromiso.objects.create(
            compromiso=commitment,
            tipo="historial",
            accion="IMPORTADO",
            descripcion=(
                f"Importado desde {filename}, hoja {sheet_name}, fila {item.row}. "
                f"Clave de origen: {item.source_key}."
            ),
            usuario="importacion_excel",
        )


class Command(BaseCommand):
    help = "Importa compromisos desde el Excel 4DX. Simula por defecto; --aplicar guarda los cambios."

    def add_arguments(self, parser):
        parser.add_argument("archivo", help="Ruta del archivo .xlsx")
        parser.add_argument("--hoja", default="Hoja1", help="Hoja que contiene los compromisos")
        parser.add_argument("--aplicar", action="store_true", help="Guarda la importación en la base de datos")
        parser.add_argument(
            "--actualizar-existentes",
            action="store_true",
            help="Sobrescribe registros importados anteriormente; sin esta opción se omiten",
        )
        parser.add_argument(
            "--reemplazar-todo",
            action="store_true",
            help="Elimina todos los compromisos y los sustituye por el Excel validado",
        )
        parser.add_argument(
            "--confirmar",
            default="",
            help="Para aplicar un reemplazo total debe indicar literalmente REEMPLAZAR",
        )

    def handle(self, *args, **options):
        path = Path(options["archivo"]).expanduser().resolve()
        if not path.is_file():
            raise CommandError(f"No existe el archivo: {path}")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise CommandError("El archivo debe tener extensión .xlsx o .xlsm")
        rows, warnings = read_rows(path, options["hoja"])
        replace_all = options["reemplazar_todo"]
        if replace_all and not rows:
            raise CommandError(
                "El Excel no contiene compromisos; el reemplazo total fue cancelado."
            )
        current_total = Compromiso.all_objects.count()
        existing = {
            row.codigo_fuente: row
            for row in Compromiso.all_objects.filter(
                codigo_fuente__in=[item.source_key for item in rows]
            )
        }
        to_create = (
            rows
            if replace_all
            else [item for item in rows if item.source_key not in existing]
        )
        to_update = (
            [] if replace_all else [item for item in rows if item.source_key in existing]
        )

        self.stdout.write(f"Archivo validado: {len(rows)} compromisos.")
        self.stdout.write(f"Nuevos: {len(to_create)}.")
        self.stdout.write(f"Ya importados: {len(to_update)}.")
        if replace_all:
            self.stdout.write(
                self.style.WARNING(
                    f"Reemplazo total: se eliminarán {current_total} compromisos actuales "
                    f"y se cargarán {len(rows)} desde el Excel."
                )
            )
        for label, row_numbers in warnings.items():
            if row_numbers:
                self.stdout.write(
                    self.style.WARNING(
                        f"{label.replace('_', ' ').capitalize()}: {len(row_numbers)} fila(s)"
                    )
                )
                if options["verbosity"] >= 2:
                    self.stdout.write("  Filas: " + ", ".join(map(str, row_numbers)))

        if not options["aplicar"]:
            self.stdout.write(
                self.style.WARNING(
                    "Simulación terminada: no se modificó la base de datos. Usa --aplicar para guardar."
                )
            )
            return
        if replace_all and options["confirmar"] != "REEMPLAZAR":
            raise CommandError(
                "Para aplicar el reemplazo total agregue --confirmar REEMPLAZAR."
            )

        created = updated = skipped = removed = 0
        try:
            with transaction.atomic():
                if replace_all:
                    removed = Compromiso.all_objects.count()
                    Compromiso.all_objects.all().delete()
                    existing = {}
                for item in rows:
                    commitment = existing.get(item.source_key)
                    if commitment and not options["actualizar_existentes"]:
                        skipped += 1
                        continue
                    values = fields_for(item, path.name, options["hoja"])
                    if commitment:
                        for field, value in values.items():
                            setattr(commitment, field, value)
                        updated += 1
                    else:
                        commitment = Compromiso(**values)
                        created += 1
                    commitment.full_clean()
                    commitment.save()
                    save_events(commitment, item, path.name, options["hoja"])
        except (ValidationError, ValueError) as error:
            raise CommandError(f"La importación fue cancelada sin guardar cambios: {error}") from error

        self.stdout.write(
            self.style.SUCCESS(
                f"Importación completa: {removed} eliminados, {created} creados, "
                f"{updated} actualizados y {skipped} omitidos."
            )
        )
